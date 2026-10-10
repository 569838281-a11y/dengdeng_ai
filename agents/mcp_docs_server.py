"""文档 MCP Server。只通过 stdio 提供表格和文档工具，不提供 JSON 名册工具。"""
import csv
import io
import logging
import os
import sys
from pathlib import Path

from mcp.server.mcpserver import MCPServer

logging.basicConfig(stream=sys.stderr, level=logging.INFO, format="%(name)s %(message)s")
logger = logging.getLogger("mcp.docs")

server = MCPServer("dengdeng-docs")

_READ_TABLE = {".csv"}
_READ_DOC = {".docx", ".pdf"}
_WRITE_DOC = {".docx", ".pdf"}


def _workspace() -> Path:
    raw = os.environ.get("MCP_WORKSPACE") or str(Path(__file__).resolve().parents[1])
    return Path(raw).resolve()


def _roots() -> tuple:
    workspace = _workspace()
    demo = (workspace / "knowledge" / "party_demo").resolve()
    user = (workspace / "data" / "party").resolve()
    user.mkdir(parents=True, exist_ok=True)
    return demo, user


def _inside(root: Path, candidate: Path) -> bool:
    try:
        candidate.relative_to(root)
        return True
    except ValueError:
        return False


def _relative(path: str) -> Path:
    text = (path or "").strip().replace("\\", "/")
    if not text or text.startswith("/") or ":" in text:
        raise ValueError("bad path")
    rel = Path(text)
    if any(part in {"..", ""} for part in rel.parts):
        raise ValueError("bad path")
    return rel


def _find_read(path: str, suffixes: set) -> Path:
    rel = _relative(path)
    if rel.suffix.lower() not in suffixes:
        raise ValueError("unsupported type")
    demo, user = _roots()
    for root in (user, demo):
        for candidate in (root / rel, root / rel.name):
            resolved = candidate.resolve()
            if _inside(root, resolved) and resolved.is_file():
                return resolved
    raise FileNotFoundError(path)


def _find_write(path: str, suffixes: set) -> Path:
    rel = _relative(path)
    if rel.suffix.lower() not in suffixes:
        raise ValueError("unsupported type")
    _demo, user = _roots()
    target = (user / rel.name).resolve()
    if not _inside(user, target):
        raise ValueError("bad path")
    target.parent.mkdir(parents=True, exist_ok=True)
    return target


def _csv_names(text: str) -> list:
    sample = (text or "").lstrip("\ufeff")
    reader = csv.DictReader(io.StringIO(sample))
    names = []
    if reader.fieldnames:
        fields = {name.strip().lower(): name for name in reader.fieldnames if name}
        key = fields.get("name")
        if key:
            for row in reader:
                name = str(row.get(key) or "").strip()
                if name and name not in names:
                    names.append(name)
            return names
    found = []
    for line in sample.splitlines():
        cell = line.split(",")[0].strip()
        if cell and cell.lower() != "name" and cell not in found:
            found.append(cell)
    return found


def _table_view(text: str) -> str:
    """名册仍按 name 行返回。第一列是 id 的宽表原样返回，四招才不会被挤掉。"""
    sample = (text or "").lstrip("\ufeff").strip()
    if not sample:
        return "empty"
    reader = csv.DictReader(io.StringIO(sample))
    fields = [name.strip() for name in (reader.fieldnames or []) if name]
    lowered = [name.lower() for name in fields]
    if lowered and lowered[0] == "id":
        return sample
    name_key = next((fields[index] for index, low in enumerate(lowered) if low == "name"), "")
    count_key = next((fields[index] for index, low in enumerate(lowered) if low == "count"), "")
    if not name_key:
        names = _csv_names(sample)
        return "\n".join(f"name: {name}" for name in names) if names else "empty"
    lines = []
    for row in reader:
        name = str(row.get(name_key) or "").strip()
        if not name:
            continue
        lines.append(f"name: {name}")
        if count_key:
            count = str(row.get(count_key) or "").strip()
            if count:
                lines.append(f"count: {count}")
    return "\n".join(lines) if lines else "empty"


@server.tool(name="read_tabular")
def read_tabular(path: str) -> str:
    """Read a CSV table. A name roster returns name lines. A sheet whose first column is id returns the table text. path example: party.csv or warehouse.csv or teams.csv."""
    target = _find_read(path, _READ_TABLE)
    text = target.read_text(encoding="utf-8-sig")
    view = _table_view(text)
    logger.info("read_tabular %s", target.name)
    return view


@server.tool(name="write_tabular")
def write_tabular(path: str, csv_text: str) -> str:
    """Write a CSV table. A name roster starts with name. A lineup sheet starts with id. csv_text includes the header."""
    target = _find_write(path, _READ_TABLE)
    body = (csv_text or "").strip().lstrip("\ufeff")
    header = body.splitlines()[0].strip().lower() if body else ""
    if not (header.startswith("name") or header.startswith("id,") or header == "id"):
        names = _csv_names("name\n" + body) if body else []
        body = "name,count\n" + "\n".join(f"{name},1" for name in names)
    target.write_text(body.rstrip() + "\n", encoding="utf-8")
    data_lines = [line for line in body.splitlines()[1:] if line.strip()]
    logger.info("write_tabular %s", target.name)
    return f"rows: {len(data_lines)}"


@server.tool(name="read_document")
def read_document(path: str) -> str:
    """Read plain text from a DOCX or PDF file. path example: notes.docx or briefing.pdf."""
    target = _find_read(path, _READ_DOC)
    suffix = target.suffix.lower()
    if suffix == ".docx":
        from docx import Document

        document = Document(str(target))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs if paragraph.text.strip())
    else:
        import pymupdf

        opened = pymupdf.open(str(target))
        try:
            text = "\n".join(page.get_text() for page in opened)
        finally:
            opened.close()
    logger.info("read_document %s", target.name)
    return (text or "").strip()[:4000] or "empty"


@server.tool(name="write_document")
def write_document(path: str, text: str, kind: str = "docx") -> str:
    """Write a DOCX or a simple PDF. kind is docx or pdf. path example: lineup.docx or lineup.pdf."""
    chosen = (kind or "").strip().lower()
    rel = _relative(path)
    if rel.suffix.lower() == ".pdf":
        chosen = "pdf"
    elif rel.suffix.lower() == ".docx":
        chosen = "docx"
    if chosen not in {"docx", "pdf"}:
        raise ValueError("kind must be docx or pdf")
    filename = rel.name
    if not filename.lower().endswith("." + chosen):
        filename = f"{Path(filename).stem}.{chosen}"
    target = _find_write(filename, _WRITE_DOC)
    payload = (text or "").strip() or "灯灯记下了这些伙伴。"
    if chosen == "docx":
        from docx import Document

        document = Document()
        for line in payload.splitlines() or [payload]:
            document.add_paragraph(line)
        document.save(str(target))
    else:
        import pymupdf

        opened = pymupdf.open()
        try:
            page = opened.new_page()
            page.insert_text((72, 72), payload[:4000], fontname="china-s", fontsize=12)
            opened.save(str(target))
        finally:
            opened.close()
    logger.info("write_document %s", target.name)
    return f"saved: {target.name}"


def main() -> None:
    server.run(transport="stdio")


if __name__ == "__main__":
    main()
